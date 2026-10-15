import importlib.util
import json
import math
import subprocess
import sys
import tempfile
import unittest
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SPEC = importlib.util.spec_from_file_location("build_examples", str(Path(__file__).with_name("build_examples.py")))
BUILD = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILD)


class SyntheticExamplesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tables = {}
        cls.events = {}
        for scenario in BUILD.SCENARIOS:
            directory = BUILD.EXAMPLES / scenario["id"]
            cls.tables[scenario["id"]] = BUILD.validate_inputs(directory)
            cls.events[scenario["id"]] = pd.read_csv(
                str(directory / "results" / "production_history.csv"), index_col=0
            )

    def test_engine_matches_published_original(self):
        committed = subprocess.check_output(["git", "show", "HEAD:sim_factory.py"], cwd=str(BUILD.ROOT))
        self.assertEqual(committed, (BUILD.ROOT / "sim_factory.py").read_bytes())

    def test_synthetic_source_is_generated_not_imported(self):
        workbook = load_workbook(str(BUILD.EXAMPLES / "synthetic_factory.xlsx"), read_only=True, data_only=True)
        try:
            for name, expected in BUILD.synthetic_data().items():
                rows = workbook[name].iter_rows(values_only=True)
                columns = next(rows)
                actual = [dict(zip(columns, row)) for row in rows]
                self.assertEqual(actual, expected)
            self.assertFalse(workbook._external_links)
        finally:
            workbook.close()

    def test_scenario_changes_are_only_declared_resources(self):
        baseline = self.tables["01_baseline"]
        capacity = self.tables["02_capacity"]
        for name in BUILD.INPUT_SHEETS:
            if name not in ["work_cells", "shifts"]:
                pd.testing.assert_frame_equal(baseline[name], capacity[name])
        for scenario in BUILD.SCENARIOS:
            tables = self.tables[scenario["id"]]
            cells = tables["work_cells"].set_index("work_cell")
            self.assertEqual(cells.loc["SYN-CELL-PACK", "number_of_equal_cells"], scenario["packing_cells"])
            opening = tables["shifts"][tables["shifts"]["shifts_A"] > 0]
            self.assertEqual(set(opening["blue_collars_per_shift"]), {scenario["workers"]})

    def test_finished_units_and_initial_stock_are_accounted_for(self):
        expected = {order["order_id"]: order["quantity"] for order in BUILD.synthetic_data()["orders"]}
        for scenario in BUILD.SCENARIOS:
            events = self.events[scenario["id"]]
            completed = events[(events["work_completion"] == "work_complete") |
                               ((events["activity"] == "take_from_warehouse") & (events["code"] == events["product"]))]
            self.assertEqual(completed.groupby("originating_order").size().to_dict(), expected)
            self.assertEqual(len(completed[completed["activity"] == "take_from_warehouse"]), 3)
            self.assertEqual(set(events["originating_order"]), set(expected))

    def test_work_respects_cells_workers_and_calendar(self):
        calendar = BUILD.synthetic_data()["calendar"]
        for scenario in BUILD.SCENARIOS:
            events = self.events[scenario["id"]]
            work = events[events["activity"].isin(["setup", "production"])]
            worker_changes = defaultdict(float)
            cell_changes = defaultdict(lambda: defaultdict(int))
            for row in work.itertuples():
                finish = (pd.Timestamp(row.time) - BUILD.START).total_seconds() / 3600
                start = finish - row.worked_time
                self.assertGreater(row.worked_time, 0)
                self.assertTrue(float(start).is_integer() and float(finish).is_integer())
                for hour in range(int(start), int(finish)):
                    day, hour_of_day = divmod(hour, 24)
                    self.assertTrue(calendar[day]["open_shifts"])
                    self.assertLess(hour_of_day, calendar[day]["shift_hours"])
                worker_changes[start] += row.crew
                worker_changes[finish] -= row.crew
                cell_changes[row.work_cell][start] += 1
                cell_changes[row.work_cell][finish] -= 1
            workers = 0
            for time in sorted(worker_changes):
                workers += worker_changes[time]
                self.assertGreaterEqual(workers, 0)
                self.assertLessEqual(workers, scenario["workers"])
            self.assertEqual(workers, 0)
            for cell, changes in cell_changes.items():
                limit = scenario["packing_cells"] if cell == "SYN-CELL-PACK" else 1
                active = 0
                for time in sorted(changes):
                    active += changes[time]
                    self.assertGreaterEqual(active, 0)
                    self.assertLessEqual(active, limit)
                self.assertEqual(active, 0)

    def test_whole_hour_constraint_is_explicit(self):
        with self.assertRaises(ValueError):
            pd.offsets.Hour(0.25)
        for events in self.events.values():
            times = pd.to_datetime(events["time"])
            self.assertTrue((times.dt.minute == 0).all())
            self.assertTrue((times.dt.second == 0).all())

    def test_complete_raw_outputs_and_hashes(self):
        for scenario in BUILD.SCENARIOS:
            directory = BUILD.EXAMPLES / scenario["id"]
            results = directory / "results"
            manifest = json.loads((results / "run_manifest.json").read_text())
            self.assertEqual(BUILD.digest(directory / "sim_factory_inputs.xlsx"), BUILD.digest(results / "sim_factory_inputs.xlsx"))
            self.assertEqual(manifest["source_engine_sha256"], BUILD.digest(BUILD.ROOT / "sim_factory.py"))
            self.assertEqual(manifest["event_log_sha256"], BUILD.digest(results / "production_history.csv"))
            self.assertFalse(manifest["engine_modified"])
            self.assertFalse(manifest["event_log_modified"])
            events = self.events[scenario["id"]]
            self.assertEqual(list(events.index), list(range(len(events))))
            self.assertEqual(set(events["activity"]), {"setup", "production", "external_lead_time", "take_from_warehouse"})
            for column in ["product", "originating_order", "code", "work_cell"]:
                self.assertTrue(events[column].dropna().str.startswith("SYN-").all())

    def test_reported_totals_match_raw_events(self):
        for scenario in BUILD.SCENARIOS:
            results = BUILD.EXAMPLES / scenario["id"] / "results"
            summary = json.loads((results / "summary.json").read_text())
            events = self.events[scenario["id"]]
            self.assertEqual(summary["events"], len(events))
            self.assertEqual(summary["activity_counts"], events["activity"].value_counts().to_dict())
            work = events[events["activity"].isin(["setup", "production"])]
            self.assertAlmostEqual(summary["worker_hours"], (work["worked_time"] * work["crew"]).sum())
            self.assertEqual(summary["simulation_stop_hours"], list(range(1, 2688, 168))[-1])
            self.assertLess(pd.to_datetime(events["time"]).max(), BUILD.START + pd.Timedelta(hours=2521))

    def test_reruns_reproduce_event_logs_byte_for_byte(self):
        for scenario in BUILD.SCENARIOS:
            directory = BUILD.EXAMPLES / scenario["id"]
            with tempfile.TemporaryDirectory(prefix="verify-synthetic-") as temporary:
                replay = Path(temporary)
                BUILD.execute(directory, replay)
                self.assertEqual(BUILD.digest(replay / "production_history.csv"),
                                 BUILD.digest(directory / "results" / "production_history.csv"))


if __name__ == "__main__":
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(SyntheticExamplesTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    BUILD.write_json(BUILD.EXAMPLES / "validation.json", {
        "status": "passed" if result.wasSuccessful() else "failed", "tests": result.testsRun,
        "failures": len(result.failures), "errors": len(result.errors),
        "validated_at_utc": datetime.utcnow().isoformat() + "Z",
        "source_engine_sha256": BUILD.digest(BUILD.ROOT / "sim_factory.py"),
        "checks": [test for test in unittest.defaultTestLoader.getTestCaseNames(SyntheticExamplesTests)],
    })
    sys.exit(0 if result.wasSuccessful() else 1)