import argparse
import csv
import hashlib
import json
import os
import platform
import random
import shutil
import subprocess
import sys
import sysconfig
import tempfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"
SEED = 41729
START = datetime(2035, 1, 1)
HORIZON_DAYS = 112
SCENARIOS = (
    {"id": "01_baseline", "label": "Constrained factory", "packing_cells": 1, "workers": 6},
    {"id": "02_capacity", "label": "Additional packing capacity", "packing_cells": 2, "workers": 8},
)
INPUT_SHEETS = (
    "production_start", "bom", "bol", "codes_descriptions", "work_cells", "initial_warehouse", "shifts"
)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_csv(path, rows, columns):
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Synthetic Factory Examples"
    workbook.properties.title = "Fictional manufacturing data - seed 41729"
    workbook.properties.description = "Entirely synthetic records; no original business records or identifier mapping."
    for name, records, columns in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for record in records:
            sheet.append([record.get(column) for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = "0F766E"
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float) or cell.data_type == "f":
                    cell.number_format = "0.000"
        for index, column in enumerate(columns, 1):
            values = [len(str(column))] + [len(str(record.get(column, ""))) for record in records[:80]]
            sheet.column_dimensions[get_column_letter(index)].width = min(48, max(14, max(values) + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))


def synthetic_data():
    generator = random.Random(SEED)
    definitions = [
        ("SYN-FG-01", "Cocoa bar lot", "finished", "PACK", 120, 2),
        ("SYN-FG-02", "Filled bite lot", "finished", "PACK", 180, 2),
        ("SYN-FG-03", "Gift selection lot", "finished", "PACK", 240, 2),
        ("SYN-M-01", "Cocoa mixture kit", "intermediate", "MIX", 60, 2),
        ("SYN-M-02", "Moulded shell kit", "intermediate", "FORM", 60, 1),
        ("SYN-M-03", "Nut filling kit", "intermediate", "MIX", 60, 2),
        ("SYN-M-04", "Plain bar kit", "intermediate", "FORM", 60, 1),
        ("SYN-M-05", "Filled bite kit", "intermediate", "FILL", 60, 2),
        ("SYN-R-01", "Cocoa ingredient kit", "purchased", "BUY", 0, 0),
        ("SYN-R-02", "Sweetener kit", "purchased", "BUY", 0, 0),
        ("SYN-R-03", "Nut ingredient kit", "purchased", "BUY", 0, 0),
        ("SYN-R-04", "Milk ingredient kit", "purchased", "BUY", 0, 0),
        ("SYN-R-05", "Wrapper kit", "purchased", "BUY", 0, 0),
        ("SYN-R-06", "Tray kit", "purchased", "BUY", 0, 0),
        ("SYN-R-07", "Presentation box kit", "purchased", "BUY", 0, 0),
    ]
    products = []
    operations = []
    for code, name, kind, stage, minutes, crew in definitions:
        products.append({"item_id": code, "name": name, "kind": kind, "model_unit": "standardised kit or lot"})
        operations.append({
            "item_id": code, "resource_id": "SYN-CELL-" + stage,
            "setup_minutes": 60 if minutes else 0,
            "production_minutes": minutes + generator.choice([0, 60]) if kind == "finished" else minutes,
            "setup_operators": crew, "production_operators": crew,
            "external_lead_days": generator.choice([1.0, 1.5, 2.0]) if kind == "purchased" else 0,
        })
    component_map = {
        "SYN-FG-01": [("SYN-M-04", 1), ("SYN-R-05", 1)],
        "SYN-FG-02": [("SYN-M-05", 1), ("SYN-R-06", 1)],
        "SYN-FG-03": [("SYN-FG-01", 1), ("SYN-FG-02", 1), ("SYN-R-07", 1)],
        "SYN-M-01": [("SYN-R-01", 2), ("SYN-R-02", 1)],
        "SYN-M-02": [("SYN-M-01", 1), ("SYN-R-04", 1)],
        "SYN-M-03": [("SYN-R-03", 1), ("SYN-R-02", 1)],
        "SYN-M-04": [("SYN-M-01", 1)],
        "SYN-M-05": [("SYN-M-02", 1), ("SYN-M-03", 1)],
    }
    components = [
        {"parent_id": parent, "component_id": child, "component_units": quantity}
        for parent, children in sorted(component_map.items()) for child, quantity in children
    ]
    resources = [
        {"resource_id": "SYN-CELL-" + stage, "department": department,
         "description": description, "parallel_cells": 1, "shift_pool": "A"}
        for stage, department, description in [
            ("MIX", "Preparation", "Mixing"), ("FORM", "Preparation", "Forming"),
            ("FILL", "Preparation", "Filling"), ("PACK", "Finishing", "Packing"),
            ("BUY", "Procurement", "Purchasing and external lead time"),
        ]
    ]
    orders = []
    for week in range(6):
        for product_index, base_quantity in enumerate([4, 3, 2], 1):
            orders.append({
                "order_id": "SYN-ORDER-{:03d}".format(len(orders) + 1),
                "customer_id": "SYN-CUSTOMER-{:02d}".format(product_index),
                "item_id": "SYN-FG-{:02d}".format(product_index),
                "quantity": base_quantity + generator.choice([-1, 0, 1]),
                "release_day": week * 7, "due_day": week * 7 + 10,
            })
    stock = [
        {"item_id": product["item_id"], "units": generator.choice([6, 8, 10])}
        for product in products if product["kind"] == "purchased"
    ] + [{"item_id": "SYN-FG-01", "units": 2}, {"item_id": "SYN-FG-02", "units": 1}]
    calendar = [
        {"simulation_day": day, "operating_date": START + timedelta(days=day),
         "open_shifts": int((START + timedelta(days=day)).weekday() < 5 and day not in [17, 38]),
         "shift_hours": 6}
        for day in range(HORIZON_DAYS + 1)
    ]
    return {"products": products, "components": components, "operations": operations,
            "resources": resources, "orders": orders, "stock": stock, "calendar": calendar}


def generate():
    data = synthetic_data()
    source_sheets = [(name, rows, list(rows[0])) for name, rows in data.items()]
    source_sheets.append(("scenarios", list(SCENARIOS), list(SCENARIOS[0])))
    write_workbook(EXAMPLES / "synthetic_factory.xlsx", source_sheets)
    for scenario in SCENARIOS:
        directory = EXAMPLES / scenario["id"]
        orders = [
            {"order_line_number": order["order_id"], "product_code": order["item_id"],
             "number_of_products": order["quantity"], "request_time": order["release_day"],
             "client": order["customer_id"], "request_date": START + timedelta(days=order["release_day"]),
             "due_date": START + timedelta(days=order["due_day"])}
            for order in data["orders"]
        ]
        bom = [{"code_produced": row["parent_id"], "code_used": row["component_id"],
                "quantity": row["component_units"]} for row in data["components"]]
        bol = [
            {"code_produced": row["item_id"], "work_cell_used": row["resource_id"],
             "setup_time": "=I{}/60".format(index), "setup_crew": row["setup_operators"],
             "production_time": "=J{}/60".format(index), "production_crew": row["production_operators"],
             "lead_time_days": row["external_lead_days"], "activity_description": "Synthetic operation",
             "setup_minutes": row["setup_minutes"], "production_minutes": row["production_minutes"]}
            for index, row in enumerate(data["operations"], 2)
        ]
        codes = [{"code_name": row["item_id"], "description": row["name"],
                  "classification": row["kind"]} for row in data["products"]]
        cells = [
            {"work_cell": row["resource_id"], "description": row["description"],
             "department": row["department"], "number_of_equal_cells":
             scenario["packing_cells"] if row["resource_id"] == "SYN-CELL-PACK" else row["parallel_cells"],
             "shift_pool": row["shift_pool"]}
            for row in data["resources"]
        ]
        warehouse = [{"code": row["item_id"], "quantity": row["units"]} for row in data["stock"]]
        shifts = []
        for day_index in range(HORIZON_DAYS + 1):
            operating = data["calendar"][HORIZON_DAYS - day_index]
            shifts.append({
                "day": day_index, "date": START + timedelta(days=day_index),
                "blue_collars_per_shift": scenario["workers"] if operating["open_shifts"] else 0,
                "shift_hours": operating["shift_hours"], "shifts_A": operating["open_shifts"],
                "simulation_day": operating["simulation_day"], "operating_date": operating["operating_date"],
            })
        model_sheets = list(zip(INPUT_SHEETS, [orders, bom, bol, codes, cells, warehouse, shifts]))
        write_workbook(directory / "sim_factory_inputs.xlsx", [(name, rows, list(rows[0])) for name, rows in model_sheets])
    print("Generated one synthetic source workbook and two seven-sheet model inputs (seed {}).".format(SEED))


def validate_inputs(directory):
    path = directory / "sim_factory_inputs.xlsx"
    workbook = load_workbook(str(path), data_only=True)
    assert workbook.sheetnames == list(INPUT_SHEETS), "Input sheet contract changed"
    assert not workbook._external_links, "External workbook link"
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "e", "Excel formula error"
    tables = pd.read_excel(str(path), sheet_name=None)
    workbook.close()
    codes = set(tables["codes_descriptions"]["code_name"])
    assert len(codes) == 15 and all(code.startswith("SYN-") for code in codes)
    assert set(tables["bol"]["code_produced"]) == codes
    assert not tables["bol"]["code_produced"].duplicated().any()
    assert set(tables["bol"]["work_cell_used"]) <= set(tables["work_cells"]["work_cell"])
    assert tables["bol"][["setup_time", "production_time"]].notnull().all().all(), "Formula caches missing; recalculate in Excel"
    assert all(abs(row.setup_time * 60 - row.setup_minutes) < 1e-8 for row in tables["bol"].itertuples())
    assert all(abs(row.production_time * 60 - row.production_minutes) < 1e-8 for row in tables["bol"].itertuples())
    assert all(float(value).is_integer() for value in tables["bol"]["setup_time"])
    assert all(float(value).is_integer() for value in tables["bol"]["production_time"])
    assert all(float(value * 24).is_integer() for value in tables["bol"]["lead_time_days"])
    graph = defaultdict(list)
    for row in tables["bom"].itertuples():
        assert row.code_produced in codes and row.code_used in codes
        assert row.quantity > 0 and row.quantity == int(row.quantity)
        graph[row.code_produced].append(row.code_used)

    def visit(code, ancestors):
        assert code not in ancestors, "Circular bill of materials"
        for child in graph[code]:
            visit(child, ancestors | {code})

    for code in codes:
        visit(code, set())
    shifts = tables["shifts"].set_index("day")
    for operating in synthetic_data()["calendar"]:
        actual = shifts.loc[HORIZON_DAYS - operating["simulation_day"]]
        assert actual["shifts_A"] == operating["open_shifts"], "Reverse calendar mapping failed"
        assert pd.Timestamp(actual["operating_date"]) == pd.Timestamp(operating["operating_date"])
    return tables


def execute(directory, destination):
    with tempfile.TemporaryDirectory(prefix="synthetic-factory-") as temporary:
        run_directory = Path(temporary)
        shutil.copyfile(str(directory / "sim_factory_inputs.xlsx"), str(run_directory / "sim_factory_inputs.xlsx"))
        environment = os.environ.copy()
        environment["PYTHONHASHSEED"] = "0"
        environment["PYTHONIOENCODING"] = "utf-8"
        environment["PYTHONWARNINGS"] = "ignore::FutureWarning,ignore::DeprecationWarning"
        completed = subprocess.run(
            [sys.executable, str(ROOT / "sim_factory.py")], cwd=str(run_directory),
            env=environment, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            universal_newlines=True, encoding="utf-8", timeout=180,
        )
        assert completed.returncode == 0, "Original engine failed: " + completed.stderr
        outputs = list((run_directory / "outputs").glob("output_*"))
        assert len(outputs) == 1, "Expected exactly one original output directory"
        assert digest(outputs[0] / "sim_factory_inputs.xlsx") == digest(directory / "sim_factory_inputs.xlsx")
        destination.mkdir(parents=True, exist_ok=True)
        for filename in ["production_history.csv", "sim_factory_inputs.xlsx"]:
            shutil.copyfile(str(outputs[0] / filename), str(destination / filename))
        (destination / "console.txt").write_text(completed.stdout + completed.stderr, encoding="utf-8")


def analyse(directory, scenario):
    tables = validate_inputs(directory)
    results = directory / "results"
    events = pd.read_csv(str(results / "production_history.csv"), index_col=0)
    events["time"] = pd.to_datetime(events["time"])
    events["worked_time"] = pd.to_numeric(events["worked_time"], errors="coerce").fillna(0)
    events["crew"] = pd.to_numeric(events["crew"], errors="coerce").fillna(0)
    completed = events[(events["work_completion"] == "work_complete") |
                       ((events["activity"] == "take_from_warehouse") & (events["code"] == events["product"]))]
    orders = []
    for order in tables["production_start"].itertuples():
        delivered = completed[completed["originating_order"] == order.order_line_number]
        assert len(delivered) <= order.number_of_products, "Duplicate final completion"
        ready = delivered["time"].max() if len(delivered) == order.number_of_products else pd.NaT
        due = pd.Timestamp(order.due_date)
        release = START + timedelta(days=int(order.request_time))
        orders.append({
            "order_id": order.order_line_number, "product": order.product_code,
            "ordered_units": int(order.number_of_products), "completed_units": len(delivered),
            "release_date": release.isoformat(), "due_date": due.isoformat(),
            "ready_date": ready.isoformat() if pd.notnull(ready) else "",
            "lead_time_days": round((ready - release).total_seconds() / 86400, 6) if pd.notnull(ready) else None,
            "late_days": round(max(0, (ready - due).total_seconds() / 86400), 6) if pd.notnull(ready) else None,
            "on_time": bool(pd.notnull(ready) and ready <= due),
        })
    assert all(row["completed_units"] == row["ordered_units"] for row in orders), "Horizon truncated demand"
    write_csv(results / "order_completion.csv", orders, list(orders[0]))
    work = events[events["activity"].isin(["setup", "production"])].copy()
    work["worker_hours"] = work["worked_time"] * work["crew"]
    weekly = work.groupby(["week", "work_cell"])[["worked_time", "worker_hours"]].sum().reset_index()
    weekly.to_csv(str(results / "weekly_workload.csv"), index=False, float_format="%.9f")
    last_time = list(range(1, HORIZON_DAYS * 24, 168))[-1]
    cells = []
    for cell in tables["work_cells"].itertuples():
        capacity = sum(max(0, min(6, last_time - day["simulation_day"] * 24)) * day["open_shifts"]
                       for day in synthetic_data()["calendar"]) * cell.number_of_equal_cells
        used = work[work["work_cell"] == cell.work_cell]["worked_time"].sum()
        cells.append({"work_cell": cell.work_cell, "parallel_cells": cell.number_of_equal_cells,
                      "worked_hours": round(float(used), 6), "available_cell_hours": capacity,
                      "horizon_utilization": round(float(used) / capacity, 6) if capacity else 0})
    write_csv(results / "cell_utilization.csv", cells, list(cells[0]))
    summary = {
        "scenario": scenario["id"], "orders": len(orders),
        "ordered_units": sum(row["ordered_units"] for row in orders),
        "completed_units": sum(row["completed_units"] for row in orders),
        "orders_on_time": sum(row["on_time"] for row in orders),
        "mean_order_lead_days": round(sum(row["lead_time_days"] for row in orders) / len(orders), 6),
        "mean_order_late_days": round(sum(row["late_days"] for row in orders) / len(orders), 6),
        "last_completion": completed["time"].max().isoformat(), "events": len(events),
        "worker_hours": round(float(work["worker_hours"].sum()), 6),
        "activity_counts": {str(key): int(value) for key, value in events["activity"].value_counts().items()},
        "simulation_stop_hours": last_time, "calendar_horizon_hours": HORIZON_DAYS * 24,
        "initial_finished_units_used": int(len(completed[completed["activity"] == "take_from_warehouse"])),
    }
    write_json(results / "summary.json", summary)
    manifest = {
        "synthetic_seed": SEED, "source_engine_sha256": digest(ROOT / "sim_factory.py"),
        "input_sha256": digest(directory / "sim_factory_inputs.xlsx"),
        "event_log_sha256": digest(results / "production_history.csv"),
        "python": platform.python_version(), "platform": sysconfig.get_platform(),
        "pandas": pd.__version__, "scenario": scenario,
        "command": "<local-python> <repository>/sim_factory.py",
        "engine_modified": False, "event_log_modified": False,
    }
    write_json(results / "run_manifest.json", manifest)
    print(json.dumps(summary, sort_keys=True))


def run():
    for scenario in SCENARIOS:
        directory = EXAMPLES / scenario["id"]
        validate_inputs(directory)
        execute(directory, directory / "results")
        analyse(directory, scenario)


def report():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as dates
    import matplotlib.pyplot as plot

    plot.rcParams.update({"font.size": 11, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.labelcolor": "#263238", "text.color": "#263238", "axes.titleweight": "bold"})
    colors = ["#B34E64", "#0F766E"]
    summaries = []
    figures = ROOT / "docs" / "images"
    figures.mkdir(parents=True, exist_ok=True)
    figure, axis = plot.subplots(figsize=(11, 5.5), constrained_layout=True)
    last = START
    for scenario, color in zip(SCENARIOS, colors):
        directory = EXAMPLES / scenario["id"] / "results"
        summary = json.loads((directory / "summary.json").read_text())
        summaries.append(summary)
        events = pd.read_csv(str(directory / "production_history.csv"), index_col=0)
        finished = events[(events["work_completion"] == "work_complete") |
                          ((events["activity"] == "take_from_warehouse") & (events["code"] == events["product"]))]
        cumulative = finished.groupby("time").size().cumsum()
        timestamps = [START] + list(pd.to_datetime(cumulative.index))
        quantities = [0] + list(cumulative)
        axis.step(timestamps, quantities, where="post", color=color, linewidth=2.4, label=scenario["label"])
        last = max(last, timestamps[-1])
    orders = synthetic_data()["orders"]
    release_days = sorted(set(row["release_day"] for row in orders))
    demand = [sum(row["quantity"] for row in orders if row["release_day"] <= day) for day in release_days]
    axis.step([START + timedelta(days=day) for day in release_days] + [last], demand + [demand[-1]],
              where="post", color="#83908D", linestyle="--", linewidth=1.8, label="Released demand")
    axis.set(title="Same 58 synthetic lots, different completion trajectories", ylabel="Finished lots ready", xlabel="Simulation date (2035)")
    axis.set_xlim(START, last + timedelta(days=3))
    axis.set_ylim(0, max(demand) + 4)
    axis.xaxis.set_major_locator(dates.WeekdayLocator(byweekday=dates.MO, interval=2))
    axis.xaxis.set_major_formatter(dates.DateFormatter("%d %b"))
    axis.grid(axis="y", color="#E3E9E6")
    axis.legend(loc="lower right", frameon=False)
    figure.savefig(str(figures / "example-completions.png"), dpi=160)
    plot.close(figure)

    figure, axes = plot.subplots(2, 1, figsize=(11, 7), sharex=True, constrained_layout=True)
    weeks = list(pd.date_range(START, last, freq="W-MON"))
    for scenario, color, axis in zip(SCENARIOS, colors, axes):
        workload = pd.read_csv(str(EXAMPLES / scenario["id"] / "results" / "weekly_workload.csv"))
        packing = workload[workload["work_cell"] == "SYN-CELL-PACK"].copy()
        packed_hours = {pd.Timestamp(row.week): row.worked_time for row in packing.itertuples()}
        capacity = [sum(day["open_shifts"] * day["shift_hours"] * scenario["packing_cells"]
                        for day in synthetic_data()["calendar"]
                        if week <= day["operating_date"] < week + timedelta(days=7)) for week in weeks]
        axis.bar(weeks, [packed_hours.get(week, 0) for week in weeks], width=4.5,
                 color=color, label="Recorded setup + production")
        axis.step(weeks, capacity, where="mid", color="#52656E", linestyle="--", label="Available packing cell-hours")
        axis.set(title=scenario["label"], ylabel="Cell-hours / week")
        axis.set_ylim(0, 84)
        axis.grid(axis="y", color="#E3E9E6")
        axis.set_axisbelow(True)
        axis.legend(loc="upper right", fontsize=9, frameon=False)
    axes[-1].xaxis.set_major_locator(dates.WeekdayLocator(byweekday=dates.MO, interval=2))
    axes[-1].xaxis.set_major_formatter(dates.DateFormatter("%d %b"))
    axes[-1].set_xlabel("Week starting (2035)")
    figure.savefig(str(figures / "example-packing.png"), dpi=160)
    plot.close(figure)
    columns = ["scenario", "ordered_units", "completed_units", "orders", "orders_on_time",
               "mean_order_lead_days", "mean_order_late_days", "last_completion", "worker_hours", "events"]
    write_csv(EXAMPLES / "comparison.csv", [{key: row[key] for key in columns} for row in summaries], columns)
    print("Rendered two measured-result charts and wrote examples/comparison.csv.")


def main():
    parser = argparse.ArgumentParser(description="Generate and run entirely synthetic examples with the original factory simulator.")
    parser.add_argument("command", choices=["generate", "run", "report"])
    arguments = parser.parse_args()
    if arguments.command == "generate":
        generate()
    elif arguments.command == "report":
        report()
    else:
        run()


if __name__ == "__main__":
    main()